"""Regression tests using actual DOCX files, not pre-marked parser strings."""
import base64
import json
from io import BytesIO
from pathlib import Path
import re
import tempfile
import unittest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from app.docx_parser import parse_docx


def mark(run, kind):
    if kind in ("FF0000", "C00000", "9C0006", "dc2626"):
        run.font.color.rgb = RGBColor.from_string(kind)
    elif kind == "underline":
        run.underline = True
    elif kind == "highlight":
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    elif kind == "red_highlight":
        run.font.highlight_color = WD_COLOR_INDEX.RED
    elif kind == "shading":
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "FF0000")
        run._r.get_or_add_rPr().append(shading)
    else:
        raise ValueError(kind)


def add_mc(doc, answer="B", kind="FF0000", target="label", inline=False):
    doc.add_paragraph("Câu 1. Chọn đáp án đúng?")
    paragraph = doc.add_paragraph() if inline else None
    for letter in "ABCD":
        p = paragraph if inline else doc.add_paragraph()
        label = p.add_run(letter)
        p.add_run(". ")
        content = p.add_run(f"Phương án {letter}")
        if letter == answer:
            mark(label if target == "label" else content, kind)
        if inline:
            p.add_run("\t")


class DocxParserTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def parse(self, doc):
        path = self.root / "exam.docx"
        doc.save(path)
        result = parse_docx(str(path), str(self.root / "images"), "/uploads")
        # Answer formatting must never leak into student-facing content.
        self.assertNotIn("[[R]]", json.dumps(result))
        self.assertNotIn("[[/R]]", json.dumps(result))
        return result

    def question(self, result):
        self.assertEqual(len(result["sections"]), 1)
        self.assertEqual(len(result["sections"][0]["questions"]), 1)
        return result["sections"][0]["questions"][0]

    def test_marking_on_label_or_content_in_paragraphs_and_inline(self):
        for kind in ("FF0000", "C00000", "9C0006", "dc2626", "underline",
                     "highlight", "red_highlight", "shading"):
            for target in ("label", "content"):
                for inline in (False, True):
                    with self.subTest(kind=kind, target=target, inline=inline):
                        doc = Document()
                        add_mc(doc, kind=kind, target=target, inline=inline)
                        result = self.parse(doc)
                        q = self.question(result)
                        self.assertEqual(q["answer"], "B")
                        self.assertEqual(q["question"], "Chọn đáp án đúng?")
                        self.assertEqual(q["options"], {c: f"Phương án {c}" for c in "ABCD"})
                        self.assertEqual(result["warnings"], [])

    def test_all_answer_letters(self):
        for answer in "ABCD":
            with self.subTest(answer=answer):
                doc = Document()
                add_mc(doc, answer=answer, kind="underline")
                self.assertEqual(self.question(self.parse(doc))["answer"], answer)

    def test_character_and_paragraph_style_inheritance(self):
        for style_type in (WD_STYLE_TYPE.CHARACTER, WD_STYLE_TYPE.PARAGRAPH):
            for kind in ("red", "highlight", "underline"):
                with self.subTest(style_type=style_type, kind=kind):
                    doc = Document()
                    base = doc.styles.add_style("AnswerBase", style_type)
                    if kind == "red":
                        base.font.color.rgb = RGBColor(192, 0, 0)
                    elif kind == "highlight":
                        base.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    else:
                        base.font.underline = True
                    derived = doc.styles.add_style("CorrectAnswer", style_type)
                    derived.base_style = base
                    add_mc(doc, answer=None)
                    option = doc.paragraphs[2]  # B
                    if style_type == WD_STYLE_TYPE.CHARACTER:
                        option.runs[0].style = derived
                    else:
                        option.style = derived
                    self.assertEqual(self.question(self.parse(doc))["answer"], "B")

    def test_explicit_formatting_off_overrides_style(self):
        doc = Document()
        style = doc.styles.add_style("Marked", WD_STYLE_TYPE.CHARACTER)
        style.font.color.rgb = RGBColor(255, 0, 0)
        style.font.underline = True
        style.font.highlight_color = WD_COLOR_INDEX.YELLOW
        add_mc(doc)
        run = doc.paragraphs[1].runs[0]  # A should not be selected
        run.style = style
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.underline = False
        highlight = OxmlElement("w:highlight")
        highlight.set(qn("w:val"), "none")
        run._r.get_or_add_rPr().append(highlight)
        self.assertEqual(self.question(self.parse(doc))["answer"], "B")

    def test_theme_color_is_resolved_instead_of_stale_rgb_fallback(self):
        doc = Document()
        add_mc(doc, answer=None)
        # The default Office theme's accent2 is red; val is only a fallback.
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        color.set(qn("w:themeColor"), "accent2")
        doc.paragraphs[2].runs[0]._r.get_or_add_rPr().append(color)
        self.assertEqual(self.question(self.parse(doc))["answer"], "B")

    def test_non_red_theme_overrides_red_rgb_fallback(self):
        doc = Document()
        add_mc(doc)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "FF0000")
        color.set(qn("w:themeColor"), "accent1")  # blue
        doc.paragraphs[1].runs[0]._r.get_or_add_rPr().append(color)
        self.assertEqual(self.question(self.parse(doc))["answer"], "B")

    def test_theme_red_tint_and_shade(self):
        for attr in ("themeTint", "themeShade"):
            with self.subTest(attr=attr):
                doc = Document()
                add_mc(doc, answer=None)
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "000000")
                color.set(qn("w:themeColor"), "accent2")
                color.set(qn(f"w:{attr}"), "99")
                doc.paragraphs[2].runs[0]._r.get_or_add_rPr().append(color)
                self.assertEqual(self.question(self.parse(doc))["answer"], "B")

    def test_option_tables_preserve_marks(self):
        for rows, cols in ((1, 4), (2, 2), (4, 1)):
            for kind in ("FF0000", "highlight", "underline"):
                with self.subTest(rows=rows, cols=cols, kind=kind):
                    doc = Document()
                    doc.add_paragraph("Câu 1. Các phương án nằm trong bảng?")
                    table = doc.add_table(rows=rows, cols=cols)
                    for letter, cell in zip("ABCD", (c for r in table.rows for c in r.cells)):
                        p = cell.paragraphs[0]
                        p.add_run(f"{letter}. ")
                        run = p.add_run(f"Phương án {letter}")
                        if letter == "C":
                            mark(run, kind)
                    result = self.parse(doc)
                    q = self.question(result)
                    self.assertEqual(q["answer"], "C")
                    self.assertEqual(q["options"]["D"], "Phương án D")
                    self.assertEqual(result["warnings"], [])

    def test_content_table_is_preserved_without_answer_markers(self):
        doc = Document()
        add_mc(doc)
        table = doc.add_table(rows=1, cols=2)
        mark(table.cell(0, 0).paragraphs[0].add_run("Dữ liệu đỏ"), "FF0000")
        table.cell(0, 1).text = "123"
        q = self.question(self.parse(doc))
        self.assertEqual(q["answer"], "B")
        payload = re.search(r"\[TABLE:([^]]+)\]", q["options"]["D"]).group(1)
        self.assertEqual(json.loads(base64.b64decode(payload)), [["Dữ liệu đỏ", "123"]])

    def test_missing_or_conflicting_markings_are_not_guessed(self):
        for conflict in (False, True):
            with self.subTest(conflict=conflict):
                doc = Document()
                add_mc(doc, answer="B" if conflict else None)
                if conflict:
                    mark(doc.paragraphs[3].runs[-1], "underline")  # C
                result = self.parse(doc)
                self.assertIsNone(self.question(result)["answer"])
                self.assertEqual(len(result["warnings"]), 1)
                self.assertIn("Nhiều phương án" if conflict else "Không tìm thấy", result["warnings"][0])

    def test_marked_question_and_whitespace_do_not_select_an_option(self):
        doc = Document()
        add_mc(doc)
        mark(doc.paragraphs[0].runs[0], "FF0000")
        mark(doc.paragraphs[1].add_run(" \t "), "underline")
        self.assertEqual(self.question(self.parse(doc))["answer"], "B")

    def test_other_font_colors_do_not_mark_an_answer(self):
        for color in ("000000", "0000FF", "008000", "ED7D31", "800080"):
            with self.subTest(color=color):
                doc = Document()
                add_mc(doc)
                doc.paragraphs[1].runs[0].font.color.rgb = RGBColor.from_string(color)
                self.assertEqual(self.question(self.parse(doc))["answer"], "B")

    def test_marked_math_content(self):
        doc = Document()
        add_mc(doc, answer=None)
        p = doc.paragraphs[2]
        p.runs[-1].text = ""
        math = OxmlElement("m:oMath")
        run = OxmlElement("m:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "FF0000")
        rpr.append(color)
        run.append(rpr)
        text = OxmlElement("m:t")
        text.text = "x=2"
        run.append(text)
        math.append(run)
        p._p.append(math)
        q = self.question(self.parse(doc))
        self.assertEqual(q["answer"], "B")
        self.assertEqual(q["options"]["B"], "[MATH:x=2]")

    def test_marked_image_content(self):
        doc = Document()
        add_mc(doc, answer=None)
        image = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="
        )
        run = doc.paragraphs[2].add_run()
        mark(run, "underline")
        run.add_picture(BytesIO(image))
        q = self.question(self.parse(doc))
        self.assertEqual(q["answer"], "B")
        self.assertIn("[IMG:/uploads/", q["options"]["B"])
        self.assertEqual(len(list((self.root / "images").iterdir())), 1)

    def test_true_false_and_short_answer_are_unchanged(self):
        doc = Document()
        doc.add_paragraph("Phần I. Trắc nghiệm nhiều lựa chọn")
        add_mc(doc, kind="highlight")
        doc.add_paragraph("Phần II. Trắc nghiệm đúng sai")
        doc.add_paragraph("Câu 1. Đánh giá các nhận định")
        for letter in "abcd":
            doc.add_paragraph(f"{letter}) Nhận định {letter}")
        table = doc.add_table(rows=1, cols=4)
        for i, letter in enumerate("abcd"):
            table.cell(0, i).text = f"{letter}) {'Đúng' if i % 2 == 0 else 'Sai'}"
        doc.add_paragraph("Phần III. Trả lời ngắn")
        doc.add_paragraph("Câu 1. Tính kết quả")
        doc.add_paragraph("Đáp án: 42")
        result = self.parse(doc)
        self.assertEqual(self.question(result)["answer"], "B")
        self.assertEqual(result["tf_questions"][0]["answers"], dict(a=True, b=False, c=True, d=False))
        self.assertEqual(result["short_answer_questions"][0]["answer"], "42")
        self.assertEqual(result["warnings"], [])


if __name__ == "__main__":
    unittest.main()
