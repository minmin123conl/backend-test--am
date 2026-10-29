"""Exercise multipart upload, persisted answers, and grading in a temporary DB."""
from contextlib import ExitStack
from io import BytesIO
import os
from pathlib import Path
import secrets
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor
from fastapi.testclient import TestClient


class DocxUploadTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        stack = ExitStack()
        cls.addClassCleanup(stack.close)
        root = Path(stack.enter_context(tempfile.TemporaryDirectory()))
        password = "Docx1!" + secrets.token_urlsafe(24)
        stack.enter_context(patch.dict(os.environ, {
            "EXAM_DB_PATH": str(root / "exam.db"),
            "EXAM_UPLOAD_DIR": str(root / "uploads"),
            "EXAM_ADMIN_USER": "docx-test-admin",
            "EXAM_ADMIN_PASSWORD": password,
            "EXAM_SECRET_KEY": secrets.token_urlsafe(32),
        }))
        # Import only after setting paths: never open the real development DB.
        from app.main import app
        from app.db import engine

        stack.callback(engine.dispose)
        cls.client = stack.enter_context(TestClient(app))
        response = cls.client.post("/api/admin/login", data={
            "username": "docx-test-admin", "password": password,
        })
        response.raise_for_status()
        cls.headers = {"Authorization": f"Bearer {response.json()['access_token']}"}

    def upload(self, doc):
        contents = BytesIO()
        doc.save(contents)
        response = self.client.post(
            "/api/admin/exams/upload",
            headers=self.headers,
            files={"file": ("marked-exam.docx", contents.getvalue(),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"title": "DOCX regression", "duration_minutes": "30"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        response = self.client.get(f"/api/admin/exams/{data['exam']['id']}", headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        return data, response.json()

    def test_marked_answers_are_persisted_and_graded(self):
        doc = Document()
        expected = ["B", "C", "D"]
        for number, answer in enumerate(expected, 1):
            doc.add_paragraph(f"Câu {number}. Chọn một đáp án")
            for letter in "ABCD":
                run = doc.add_paragraph().add_run(f"{letter}. Nội dung {letter}")
                if letter == answer:
                    if number == 1:
                        run.font.color.rgb = RGBColor(192, 0, 0)
                    elif number == 2:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    else:
                        run.underline = True
        upload, exam = self.upload(doc)
        self.assertEqual(upload["stats"], dict(num_mc=3, num_tf=0, num_sa=0, missing_answers=0))
        self.assertEqual(upload["warnings"], [])
        self.assertEqual([q["data"]["answer"] for q in exam["questions"]], expected)

        from app.db import SessionLocal
        from app.models import Exam
        from app.scoring import grade_attempt

        with SessionLocal() as db:
            stored = db.get(Exam, exam["id"])
            answers = {str(q["id"]): answer for q, answer in zip(exam["questions"], expected)}
            score, total, correct, _ = grade_attempt(stored, answers)
            self.assertEqual((score, total, correct), (3.0, 3.0, 3))

    def test_table_options_and_existing_question_types(self):
        doc = Document()
        doc.add_paragraph("Phần I. Trắc nghiệm nhiều lựa chọn")
        doc.add_paragraph("Câu 1. Chọn một phương án trong bảng")
        table = doc.add_table(rows=2, cols=2)
        for letter, cell in zip("ABCD", (c for row in table.rows for c in row.cells)):
            run = cell.paragraphs[0].add_run(f"{letter}. Nội dung {letter}")
            if letter == "C":
                run.underline = True
        doc.add_paragraph("Phần II. Trắc nghiệm đúng sai")
        doc.add_paragraph("Câu 1. Đánh giá nhận định")
        for letter in "abcd":
            doc.add_paragraph(f"{letter}) Nhận định {letter}")
        table = doc.add_table(rows=1, cols=4)
        for i, letter in enumerate("abcd"):
            table.cell(0, i).text = f"{letter}) {'Đúng' if i % 2 == 0 else 'Sai'}"
        doc.add_paragraph("Phần III. Trả lời ngắn")
        doc.add_paragraph("Câu 1. Tính kết quả")
        doc.add_paragraph("Đáp án: 42")
        upload, exam = self.upload(doc)
        self.assertEqual(upload["stats"], dict(num_mc=1, num_tf=1, num_sa=1, missing_answers=0))
        self.assertEqual(upload["warnings"], [])
        mc, tf, sa = exam["questions"]
        self.assertEqual(mc["data"]["answer"], "C")
        self.assertEqual(tf["data"]["answers"], dict(a=True, b=False, c=True, d=False))
        self.assertEqual(sa["data"]["answer"], "42")

    def test_conflicting_marks_surface_as_missing_answer_and_warning(self):
        doc = Document()
        doc.add_paragraph("Câu 1. Chọn một phương án")
        for letter in "ABCD":
            run = doc.add_paragraph().add_run(f"{letter}. Nội dung {letter}")
            if letter in "BC":
                run.underline = True
        upload, exam = self.upload(doc)
        self.assertEqual(upload["stats"]["missing_answers"], 1)
        self.assertIsNone(exam["questions"][0]["data"]["answer"])
        self.assertIn("Nhiều phương án", upload["warnings"][0])


if __name__ == "__main__":
    unittest.main()
